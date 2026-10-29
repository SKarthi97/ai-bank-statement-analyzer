"""Regenerate the synthetic XLSX fixtures: ``python tests/fixtures/build_xlsx_fixtures.py``.

All data is invented. The fixtures are committed so tests do not depend on this script running.
"""

from __future__ import annotations

import zipfile
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook

HERE = Path(__file__).parent

# Same transactions as deposit_basic.csv, so the same hand-computed oracle applies.
BASIC = [
    (date(2024, 1, 2), "SALARY ACME CORP", 3000.00, 3500.00),
    (date(2024, 1, 3), "GROCERY MART #1234 CARD 5678", -85.40, 3414.60),
    (date(2024, 1, 15), "CITY POWER AND LIGHT", -120.00, 3294.60),
    (date(2024, 1, 31), "NETFLIX.COM", -15.99, 3278.61),
    (date(2024, 2, 1), "ATM WITHDRAWAL MAIN ST", -200.00, 3078.61),
    (date(2024, 2, 5), "GROCERY MART #1234 CARD 5678", -92.10, 2986.51),
    (date(2024, 2, 5), "GROCERY MART #1234 CARD 5678", -92.10, 2894.41),
    (date(2024, 2, 10), "REFUND GROCERY MART", 12.50, 2906.91),
    (date(2024, 2, 29), "SALARY ACME CORP", 3000.00, 5906.91),
]


def bank_basic() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Transactions"
    ws.append(["Statement for Checking (synthetic)"])
    ws.append(["Period: Jan-Feb 2024"])
    ws.append([])
    ws.append(["Date", "Description", "Amount", "Balance"])
    for row in BASIC:
        ws.append(list(row))
    ws.append([None, "Closing Balance", None, 5906.91])
    wb.save(HERE / "bank_basic.xlsx")


def multi_sheet() -> None:
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    summary.append(["Annual overview"])
    summary.append(["Nothing tabular here"])
    hidden = wb.create_sheet("Hidden")
    hidden.append(["Date", "Description", "Amount"])
    hidden.append([date(2020, 1, 1), "SHOULD NOT BE USED", -1.0])
    hidden.sheet_state = "hidden"
    ws = wb.create_sheet("Transactions")
    ws.append(["Date", "Description", "Amount", "Balance"])
    ws.append([date(2024, 1, 2), "SALARY", 1000.50, 1000.50])
    ws.append([date(2024, 1, 3), "COFFEE", -3.25, 997.25])
    wb.save(HERE / "multi_sheet.xlsx")


def text_amounts_eu() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Umsaetze"
    ws.append(["Datum", "Beschreibung", "Soll", "Haben", "Saldo"])
    ws.append(["02.01.2024", "GEHALT FIRMA GMBH", None, "2.500,50", "2.500,50"])
    ws.append(["05.01.2024", "MIETE JANUAR", "1.200,00", None, "1.300,50"])
    ws.append(["25.01.2024", "SUPERMARKT", "45,30", None, "1.255,20"])
    wb.save(HERE / "text_amounts_eu.xlsx")


def float_noise() -> None:
    wb = Workbook()
    ws = wb.active
    ws.append(["Date", "Description", "Amount"])
    ws.append([date(2024, 1, 2), "CLEAN", -10.25])
    ws.append([date(2024, 1, 3), "BINARY FLOAT NOISE", 0.3])
    target = HERE / "float_noise.xlsx"
    wb.save(target)
    # openpyxl trims floats when saving, but Excel stores computed values with 17 significant
    # digits (=0.1+0.2 is stored as 0.30000000000000004). Patch the XML to match real Excel.
    with zipfile.ZipFile(target) as src:
        parts = {i.filename: src.read(i.filename) for i in src.infolist()}
    sheet = parts["xl/worksheets/sheet1.xml"]
    assert b"<v>0.3</v>" in sheet
    parts["xl/worksheets/sheet1.xml"] = sheet.replace(b"<v>0.3</v>", b"<v>0.30000000000000004</v>")
    with zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED) as out:
        for name, data in parts.items():
            out.writestr(name, data)


def datetime_with_time() -> None:
    wb = Workbook()
    ws = wb.active
    ws.append(["Date", "Description", "Amount"])
    ws.append([datetime(2024, 1, 2, 14, 30), "AFTERNOON PURCHASE", -5.00])
    wb.save(HERE / "datetime_with_time.xlsx")


if __name__ == "__main__":
    for build in (bank_basic, multi_sheet, text_amounts_eu, float_noise, datetime_with_time):
        build()
    print("fixtures written to", HERE)
