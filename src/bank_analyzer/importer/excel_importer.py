"""XLSX extraction with openpyxl in read-only, values-only mode.

Safety: macros are never executed (and macro workbooks are rejected earlier), external links are
not followed (``keep_links=False``), and only cached cell values are read. Numeric cells are
converted to text from the shortest round-trip representation of the stored double, so ``85.4``
becomes "85.4". A cell holding binary-float noise such as 0.30000000000000004 is passed on as
is and is then rejected by the amount parser instead of being silently rounded.
"""

from __future__ import annotations

import math
from datetime import date, datetime, time
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from bank_analyzer.domain.errors import ImportStructureError
from bank_analyzer.domain.imports import RawRow, RawTable
from bank_analyzer.domain.models import FileKind
from bank_analyzer.importer.column_mapper import find_header_row

MAX_ROWS = 200_000
MAX_COLUMNS = 200


class _Counters:
    def __init__(self) -> None:
        self.datetimes_with_time = 0
        self.non_finite = 0


def _to_text(value: Any, counters: _Counters) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if not math.isfinite(value):
            counters.non_finite += 1
            return ""
        return format(Decimal(repr(value)), "f")
    if isinstance(value, Decimal):
        return format(value, "f")
    if isinstance(value, datetime):
        if value.time() != time(0, 0):
            counters.datetimes_with_time += 1
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return " ".join(str(value).split())


def list_sheets(path: Path) -> list[str]:
    """Names of the visible sheets, in workbook order."""
    workbook = _open(path)
    try:
        return [ws.title for ws in workbook.worksheets if ws.sheet_state == "visible"]
    finally:
        workbook.close()


def _open(path: Path) -> Any:
    try:
        return load_workbook(path, read_only=True, data_only=True, keep_links=False)
    except Exception as exc:  # third-party parser boundary: any failure means unreadable
        raise ImportStructureError("The workbook could not be read.") from exc


def _read_sheet(
    sheet: Any, header_row: int | None, max_rows: int, counters: _Counters
) -> tuple[tuple[str, ...], list[RawRow], list[str]] | None:
    records: list[tuple[int, list[str]]] = []
    for number, values in enumerate(sheet.iter_rows(values_only=True), start=1):
        cells = [_to_text(v, counters) for v in values[:MAX_COLUMNS]]
        while cells and not cells[-1]:
            cells.pop()
        records.append((number, cells))
        if len(records) > max_rows + 1000:
            raise ImportStructureError(f"The sheet has more than {max_rows:,} rows.")

    if header_row is not None:
        index: int | None = next((i for i, (n, _) in enumerate(records) if n == header_row), None)
    else:
        index = find_header_row([c for _, c in records])
    if index is None:
        return None

    warnings: list[str] = []
    if records[index][0] > 1:
        warnings.append(f"Skipped {records[index][0] - 1} row(s) before the header row.")
    headers = tuple(records[index][1])
    rows: list[RawRow] = []
    blank = 0
    for number, cells in records[index + 1 :]:
        if not any(cells):
            blank += 1
            continue
        if len(rows) >= max_rows:
            raise ImportStructureError(f"The sheet has more than {max_rows:,} rows.")
        rows.append(RawRow(row_index=number, cells=tuple(cells), location=sheet.title))
    if blank:
        warnings.append(f"Ignored {blank} blank row(s).")
    return headers, rows, warnings


def read_xlsx(
    path: Path,
    *,
    sha256: str,
    sheet: str | None = None,
    header_row: int | None = None,
    max_rows: int = MAX_ROWS,
) -> RawTable:
    """Extract one sheet of a workbook into a RawTable.

    Without ``sheet`` the first visible sheet that contains a recognisable header is used, and
    the choice is reported in the warnings. Date cells are returned as ISO ``YYYY-MM-DD`` text.
    """
    workbook = _open(path)
    try:
        visible = [ws for ws in workbook.worksheets if ws.sheet_state == "visible"]
        if sheet is not None:
            candidates = [ws for ws in visible if ws.title == sheet]
            if not candidates:
                raise ImportStructureError("The chosen sheet does not exist or is hidden.")
        else:
            candidates = visible
        counters = _Counters()
        for ws in candidates:
            result = _read_sheet(ws, header_row, max_rows, counters)
            if result is None:
                continue
            headers, rows, warnings = result
            notes = [f"Sheet '{ws.title}' was used."]
            if sheet is None and len(visible) > 1:
                notes.append(
                    f"The workbook has {len(visible)} visible sheets; verify this is the right one."
                )
            if counters.datetimes_with_time:
                notes.append(
                    f"{counters.datetimes_with_time} date cell(s) had a time part, which was dropped."
                )
            if counters.non_finite:
                notes.append(
                    f"{counters.non_finite} cell(s) held NaN or infinity and were left empty."
                )
            notes.append("Date cells were converted to ISO format (YYYY-MM-DD).")
            return RawTable(
                file_sha256=sha256,
                file_kind=FileKind.XLSX,
                headers=headers,
                rows=tuple(rows),
                warnings=tuple(notes + warnings),
            )
        raise ImportStructureError(
            "No header row was recognised. Choose the sheet or header row manually."
        )
    finally:
        workbook.close()
